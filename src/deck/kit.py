"""Палитра сайта «Агрокаст» и мелкие помощники для python-pptx: текст, карточки, диаграммы, картинки."""
from __future__ import annotations

import re
import zipfile
from pathlib import Path

from lxml import etree
from pptx.chart.data import CategoryChartData, XyChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import (XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION, XL_MARKER_STYLE, XL_TICK_LABEL_POSITION,
                             XL_TICK_MARK)
from pptx.enum.dml import MSO_LINE
from pptx.enum.lang import MSO_LANGUAGE_ID
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

# ---- палитра: токены web/css/app.css (тёмно-зелёный + тёплый жёлтый), без цветов и логотипа банка
BG = "F6F4EE"          # фон слайдов — как у сайта
SURFACE = "FFFFFF"
BORDER = "E6E1D6"
INK = "1E2420"
INK2 = "5A615B"
INK3 = "8C918A"
BRAND = "1F4D3A"       # главный тёмно-зелёный
BRAND_DK = "143629"    # фон тёмных слайдов
BRAND2 = "2E6B4F"
BRAND_SOFT = "E3EEE6"
GOLD = "E2B04A"
GOLD2 = "C8962E"
GOLD_SOFT = "F8EDD3"
POS = "1B7F6B"
NEG = "A6611A"
# порядковая шкала «лесенки»: тёплый серый → тёмно-зелёный (светлее 250 по яркости не берём)
RAMP5 = ["C9C4B5", "A9C6B3", "6FA286", "2E6B4F", "143629"]
# цвета культур — из web/data/crops.json
FONT = "Calibri"
SW, SH = 13.333, 7.5
MX = 0.6  # поля слева и справа


def rgb(h: str) -> RGBColor:
    return RGBColor.from_string(h)


# ================================================================ примитивы

def _clear_effects(shape) -> None:
    """Убирает эффекты и добавляет пустой effectLst: иначе фигура берёт тень из стиля темы."""
    spPr = shape._element.spPr
    for e in spPr.findall(qn("a:effectLst")):
        spPr.remove(e)
    etree.SubElement(spPr, qn("a:effectLst"))


def shadow(shape, blur: float = 16, dist: float = 4, alpha: int = 14, color: str = INK) -> None:
    _clear_effects(shape)
    spPr = shape._element.spPr
    eff = spPr.find(qn("a:effectLst"))
    sh = etree.SubElement(eff, qn("a:outerShdw"), blurRad=str(int(blur * 12700)), dist=str(int(dist * 12700)),
                          dir="5400000", algn="t", rotWithShape="0")
    clr = etree.SubElement(sh, qn("a:srgbClr"), val=color)
    etree.SubElement(clr, qn("a:alpha"), val=str(alpha * 1000))


def set_alpha(shape, pct_opaque: int) -> None:
    """Прозрачность сплошной заливки: pct_opaque — доля непрозрачности, %."""
    sf = shape._element.spPr.find(qn("a:solidFill"))
    if sf is not None and len(sf):
        clr = sf[0]
        for a in clr.findall(qn("a:alpha")):
            clr.remove(a)
        etree.SubElement(clr, qn("a:alpha"), val=str(pct_opaque * 1000))


def rect(slide, x, y, w, h, fill=SURFACE, line=None, line_w=0.75, radius=0.0, shape=None, shade=False,
         name=None, dash=None):
    kind = shape or (MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE)
    s = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    if radius and kind == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = min(0.5, radius / min(w, h))
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = rgb(fill)
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = rgb(line)
        s.line.width = Pt(line_w)
        if dash:
            s.line.dash_style = dash
    _clear_effects(s)
    if shade:
        shadow(s)
    if name:
        s.name = name
    s.text_frame.text = ""
    return s


def _fill_paragraphs(tf, paras, size, color, bold, align, font, line_spacing, space_after, italic=False):
    """paras: str | [para]; para: str | [run]; run: str | (text, {size,color,bold,italic})."""
    if isinstance(paras, str):
        paras = [paras]
    first = True
    for para in paras:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[align]
        if line_spacing:
            p.line_spacing = line_spacing
        if space_after is not None:
            p.space_after = Pt(space_after)
        runs = [para] if isinstance(para, (str, tuple)) else para
        for r in runs:
            text, o = (r, {}) if isinstance(r, str) else r
            run = p.add_run()
            run.text = text
            f = run.font
            f.name = o.get("font", font)
            f.size = Pt(o.get("size", size))
            f.bold = o.get("bold", bold)
            f.italic = o.get("italic", italic)
            f.color.rgb = rgb(o.get("color", color))
            f.language_id = MSO_LANGUAGE_ID.RUSSIAN


def text(slide, x, y, w, h, paras, size=14, color=INK, bold=False, align="l", anchor="t", font=FONT,
         line_spacing=None, space_after=None, name=None, italic=False, margin=0.0):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    for side in ("left", "right", "top", "bottom"):
        setattr(tf, f"margin_{side}", Inches(margin))
    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[anchor]
    _fill_paragraphs(tf, paras, size, color, bold, align, font, line_spacing, space_after, italic)
    if name:
        tb.name = name
    return tb


def shape_text(shape, paras, size=14, color=INK, bold=False, align="c", anchor="m", margin=0.08, font=FONT,
               line_spacing=None, space_after=None):
    tf = shape.text_frame
    tf.word_wrap = True
    for side in ("left", "right", "top", "bottom"):
        setattr(tf, f"margin_{side}", Inches(margin))
    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[anchor]
    # убрать пустой первый абзац
    for p in list(tf.paragraphs)[1:]:
        p._p.getparent().remove(p._p)
    for r in list(tf.paragraphs[0].runs):
        r._r.getparent().remove(r._r)
    _fill_paragraphs(tf, paras, size, color, bold, align, font, line_spacing, space_after)


def pill(slide, x, y, w, h, label, fill=BRAND_SOFT, color=BRAND, size=12, bold=True, line=None):
    s = rect(slide, x, y, w, h, fill=fill, line=line, radius=h / 2)
    shape_text(s, label, size=size, color=color, bold=bold, margin=0.05)
    return s


def card(slide, x, y, w, h, fill=SURFACE, radius=0.14, shade=True, line=BORDER):
    return rect(slide, x, y, w, h, fill=fill, line=line if fill == SURFACE else None, line_w=0.75, radius=radius, shade=shade)


def circle(slide, x, y, d, fill=BRAND, label="", color="FFFFFF", size=16, bold=True):
    s = rect(slide, x, y, d, d, fill=fill, shape=MSO_SHAPE.OVAL)
    if label:
        shape_text(s, label, size=size, color=color, bold=bold, margin=0)
    return s


def line(slide, x1, y1, x2, y2, color=BORDER, width=1.0, dash=None):
    c = slide.shapes.add_connector(1, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = rgb(color)
    c.line.width = Pt(width)
    if dash:
        c.line.dash_style = dash
    return c


def slide_number(slide, x, y, w=0.6, size=10, color=INK3):
    """Поле «номер слайда» — не теряется при перестановке."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(0.25))
    tf = tb.text_frame
    for side in ("left", "right", "top", "bottom"):
        setattr(tf, f"margin_{side}", 0)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    fld = etree.SubElement(p._p, qn("a:fld"), id="{B6F15528-21DE-4FAA-801E-634DDDAF4B2B}", type="slidenum")
    rPr = etree.SubElement(fld, qn("a:rPr"), lang="ru-RU", sz=str(size * 100))
    sf = etree.SubElement(rPr, qn("a:solidFill"))
    etree.SubElement(sf, qn("a:srgbClr"), val=color)
    etree.SubElement(rPr, qn("a:latin"), typeface=FONT)
    t = etree.SubElement(fld, qn("a:t"))
    t.text = "‹#›"
    tb.name = "Номер слайда"
    return tb


def alt(shape, descr: str) -> None:
    """Альтернативный текст для экранных дикторов."""
    shape._element.xpath(".//p:cNvPr")[0].set("descr", descr)


def picture(slide, path, x, y, w=None, h=None, radius_in=0.12, shade=True, border=BORDER, descr=None):
    pic = slide.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w) if w else None, Inches(h) if h else None)
    if descr:
        alt(pic, descr)
    pic.auto_shape_type = MSO_SHAPE.ROUNDED_RECTANGLE
    gd = pic._element.spPr.find(qn("a:prstGeom")).find(qn("a:avLst"))
    for g in list(gd):
        gd.remove(g)
    short = min(pic.width, pic.height)
    etree.SubElement(gd, qn("a:gd"), name="adj", fmla=f"val {int(min(0.5, Inches(radius_in) / short) * 100000)}")
    if border:
        pic.line.color.rgb = rgb(border)
        pic.line.width = Pt(0.75)
    if shade:
        shadow(pic, blur=22, dist=6, alpha=20)
    return pic


# ================================================================ диаграммы

def _font(f, size, color=INK2, bold=False):
    f.size = Pt(size)
    f.name = FONT
    f.bold = bold
    f.color.rgb = rgb(color)


def style_axes(chart, cat_size=12, val_size=11, show_val_axis=True, grid=True, vmin=None, vmax=None, major=None,
               val_fmt=None, cat_color=INK2):
    ca, va = chart.category_axis, chart.value_axis
    _font(ca.tick_labels.font, cat_size, cat_color)
    ca.format.line.color.rgb = rgb(BORDER)
    ca.has_major_gridlines = False
    ca.major_tick_mark = XL_TICK_MARK.NONE
    va.visible = show_val_axis
    va.has_major_gridlines = grid and show_val_axis
    if va.has_major_gridlines:
        va.major_gridlines.format.line.color.rgb = rgb(BORDER)
        va.major_gridlines.format.line.width = Pt(0.75)
    va.format.line.fill.background()
    va.major_tick_mark = XL_TICK_MARK.NONE
    _font(va.tick_labels.font, val_size, INK3)
    if vmin is not None:
        va.minimum_scale = vmin
    if vmax is not None:
        va.maximum_scale = vmax
    if major is not None:
        va.major_unit = major
    if val_fmt:
        va.tick_labels.number_format = val_fmt
        va.tick_labels.number_format_is_linked = False


def add_chart(slide, kind, data, x, y, w, h, name=None):
    gf = slide.shapes.add_chart(kind, Inches(x), Inches(y), Inches(w), Inches(h), data)
    ch = gf.chart
    ch.font.name = FONT
    ch.font.size = Pt(12)
    ch.font.color.rgb = rgb(INK2)
    ch.has_title = False
    if name:
        gf.name = name
        alt(gf, "Диаграмма: " + name)
    return ch


def legend(chart, pos=XL_LEGEND_POSITION.TOP, size=12):
    chart.has_legend = True
    chart.legend.position = pos
    chart.legend.include_in_layout = False
    _font(chart.legend.font, size, INK2)


def labels(series_or_plot, fmt, size=13, color=INK, bold=True, pos=None):
    if hasattr(series_or_plot, "has_data_labels"):  # plot: сначала включить подписи
        series_or_plot.has_data_labels = True
    dl = series_or_plot.data_labels
    dl.show_value = True
    dl.number_format = fmt
    dl.number_format_is_linked = False
    _font(dl.font, size, color, bold)
    if pos is not None:
        dl.position = pos
    return dl


def color_points(series, colors):
    for i, c in enumerate(colors):
        pt = series.points[i]
        pt.format.fill.solid()
        pt.format.fill.fore_color.rgb = rgb(c)
        pt.format.line.fill.background()


# ================================================================ тема файла

def patch_theme(path: Path) -> None:
    """Подменяет цвета и шрифты темы презентации на палитру проекта (диаграммы и новые фигуры наследуют её)."""
    colors = {"dk2": INK, "lt2": BG, "accent1": BRAND, "accent2": GOLD, "accent3": BRAND2, "accent4": GOLD2,
              "accent5": "6FA286", "accent6": INK2, "hlink": BRAND2, "folHlink": INK2}
    tmp = Path(str(path) + ".tmp")
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if re.fullmatch(r"ppt/theme/theme\d+\.xml", item.filename):
                s = data.decode("utf-8")
                for k, v in colors.items():
                    s = re.sub(rf"(<a:{k}>)\s*<a:(?:srgbClr|sysClr)[^>]*?/>\s*(</a:{k}>)", rf'\1<a:srgbClr val="{v}"/>\2', s)
                s = re.sub(r'(<a:(?:major|minor)Font>\s*<a:latin typeface=")[^"]*(")', rf"\g<1>{FONT}\2", s)
                data = s.encode("utf-8")
            zout.writestr(item, data)
    tmp.replace(path)
