"""Презентация для жюри: .pptx + PDF + сценарий. Все числа берутся из таблиц и отчётов проекта.

    uv run python -I -X utf8 src/deck/build_deck.py            # pptx, script.md, PDF через PowerPoint
    uv run python -I -X utf8 src/deck/build_deck.py --no-pdf   # без PDF
    uv run python -I -X utf8 src/deck/build_deck.py --png      # плюс PNG слайдов в deck/build/png (для проверки глазами)

PDF и PNG делает установленный PowerPoint (COM), без него — только .pptx.
"""
from __future__ import annotations

import argparse
import shutil
import subprocess
import sys
import tomllib
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "src"))

from pptx import Presentation  # noqa: E402
from pptx.util import Emu  # noqa: E402

from deck import facts, kit, script, slides  # noqa: E402

DECK = ROOT / "deck"
OUT = DECK / "out"
NAME = "agrokast_presentation"

PS_EXPORT = """
$ErrorActionPreference = 'Stop'
$app = New-Object -ComObject PowerPoint.Application
try {{
    $pres = $app.Presentations.Open('{pptx}', $true, $false, $false)   # только чтение, без окна
    {pdf_cmd}
    {png_cmd}
    $pres.Close()
}} finally {{
    if ($app.Presentations.Count -eq 0) {{ $app.Quit() }}              # чужие открытые презентации не трогаем
}}
"""


def export(pptx: Path, pdf: Path | None, png_dir: Path | None) -> bool:
    pdf_cmd = f"$pres.SaveAs('{pdf}', 32)" if pdf else ""          # 32 = ppSaveAsPDF
    if png_dir:
        shutil.rmtree(png_dir, ignore_errors=True)
        png_dir.mkdir(parents=True, exist_ok=True)
        png_cmd = f"$pres.Export('{png_dir}', 'PNG', 1920, 1080)"
    else:
        png_cmd = ""
    ps = PS_EXPORT.format(pptx=pptx, pdf_cmd=pdf_cmd, png_cmd=png_cmd)
    r = subprocess.run(["powershell", "-NoProfile", "-NonInteractive", "-Command", ps], capture_output=True, text=True, timeout=300)
    if r.returncode != 0:
        print("[export] PowerPoint:", (r.stderr or r.stdout).strip()[:600], file=sys.stderr)
        return False
    return True


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--no-pdf", action="store_true")
    ap.add_argument("--png", action="store_true", help="выгрузить PNG слайдов в deck/build/png")
    a = ap.parse_args()

    cfg = tomllib.loads((DECK / "config.toml").read_text(encoding="utf-8"))
    F = facts.load()
    for m in F["missing"]:
        print("[deck] нет данных:", m, file=sys.stderr)

    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)
    for fn in slides.SLIDES:
        fn(prs, F, cfg)

    sc = script.build(F, cfg)
    for sld, item in zip(prs.slides, sc):
        sld.notes_slide.notes_text_frame.text = item["text"]
    cp = prs.core_properties
    cp.title = "Агрокаст: прогноз урожайности и оптимизация структуры посевов"
    cp.author = cfg["author"]
    cp.subject = "Кейс «Своё Фермерство», задача 5"
    cp.language = "ru-RU"

    OUT.mkdir(parents=True, exist_ok=True)
    pptx = OUT / f"{NAME}.pptx"
    prs.save(pptx)
    kit.patch_theme(pptx)
    (DECK / "script.md").write_text(script.render(F, cfg, sc), encoding="utf-8")
    print(f"[deck] {pptx.relative_to(ROOT)} — {len(prs.slides)} слайдов; сценарий deck/script.md")

    if not a.no_pdf or a.png:
        pdf = None if a.no_pdf else OUT / f"{NAME}.pdf"
        png = DECK / "build" / "png" if a.png else None
        if export(pptx, pdf, png):
            done = [str(pdf.relative_to(ROOT))] if pdf else []
            print("[deck] экспорт PowerPoint:", ", ".join(done + (["PNG в deck/build/png"] if png else [])))
    o = F.get("opt")
    if o and not o["fresh"]:
        print("[deck] ВНИМАНИЕ: оптимизация посчитана не на итоговой модели или старше её прогнозов", file=sys.stderr)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
