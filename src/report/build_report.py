"""Сборка научного отчёта: report/text/*.md + таблицы и графики проекта → report/out/*.docx и *.pdf.

Цифры в тексте не пишутся руками: в markdown стоят ключи {{key}}, а значения считаются здесь из
reports/tables/*, data/model/* и инвентаря данных. Если результата ещё нет (ночные расчёты),
на его месте появляется жёлтая пометка «ждёт данных», и отчёт всё равно собирается.

    uv run --group report python -I -X utf8 src/report/build_report.py            # всё
    uv run --group report python -I -X utf8 src/report/build_report.py --no-pdf   # только .docx
    uv run --group report python -I -X utf8 src/report/build_report.py --stats --shap  # пересчитать инвентарь данных и SHAP

Разметка в report/text/*.md: заголовки #/##/###, абзацы, списки «- » и «1. », таблицы markdown,
**жирный**, *курсив*, `код`, x_{индекс} и x^{степень}, формулы `$$ … $$` отдельной строкой и директивы:
`::figure id | подпись`, `::image путь | подпись`, `::table id | подпись`, `::keyfigures`, `::pending текст`,
`::if флаг` / `::else` / `::endif`, `::pagebreak`.
"""
from __future__ import annotations

import argparse
import re
import sys
import time
import tomllib
from pathlib import Path

import numpy as np
import pandas as pd

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "src"))

from report import data_stats, figures, sources as S, style  # noqa: E402
from report.style import CROP_NAME, CROP_ORDER, CROP_SHORT  # noqa: E402

TEXT = ROOT / "report" / "text"
OUT = ROOT / "report" / "out"
FIG = S.BUILD / "fig"
NAME = "rshb_yield_forecast_report"
PENDING = "ждёт данных"


# ================================================================ цифры

def fnum(x, d=0):
    return S.fmt(x, d)


def compute_numbers() -> tuple[dict, dict]:
    """Ключи {{…}} для текста и флаги для ::if. Каждый блок считается отдельно: сбой одного не роняет отчёт."""
    N: dict[str, str] = {"date_build": time.strftime("%d.%m.%Y")}
    F: dict[str, bool] = {}

    def block(fn):
        try:
            fn()
        except Exception as e:  # noqa: BLE001
            print(f"[numbers] {fn.__name__}: {e!r}", file=sys.stderr)

    def data():
        st = data_stats.load()
        raw = st["layers"]["raw"]
        inv = S.emiss_inventory()
        srcs = [k for k, v in st["raw_sources"].items() if v["bytes"] > 0]
        N["n_sources"] = str(len(srcs))
        N["raw_gb"] = fnum(raw["bytes"] / 2**30, 1)
        N["raw_files"] = fnum(raw["files"])
        N["raw_rows_mln"] = fnum(raw["rows"] / 1e6, 1)
        N["emiss_ok"] = str(inv.get("n_ok", 74))
        N["emiss_obs_mln"] = fnum(inv.get("rows", 0) / 1e6, 1)
        N["total_rows_mln"] = fnum((raw["rows"] + inv.get("rows", 0)) / 1e6, 0)
        ct = st["clean_tables"]
        N["clean_tables"] = str(len(ct))
        mt = ct.get("model_table", {})
        N["model_rows"] = fnum(mt.get("rows"))
        N["model_cols"] = fnum(mt.get("cols"))
        N["n_tests"] = str(S.tests_count())

    def sweep():
        c = S.n_configs()
        N["cfg_pre"], N["cfg_v1"], N["cfg_v2"] = fnum(c["pre"]), fnum(c["v1"]), fnum(c["v2"])
        N["cfg_total"] = fnum(c["total"])
        N["optuna_trials"], N["optuna_studies"] = fnum(c["optuna_trials"]), fnum(c["optuna_studies"])
        r = S.sweeps()
        b = S.best(r[r.ck.isin(CROP_ORDER)])
        N["gain_bl_min"], N["gain_bl_max"] = fnum(b.gain_pct.min()), fnum(b.gain_pct.max())
        N["gain_naive_min"], N["gain_naive_max"] = fnum(b.gain_naive.min()), fnum(b.gain_naive.max())
        N["gain_naive_mean"] = fnum(b.gain_naive.mean())
        for k in CROP_ORDER:
            if k in b.index:
                N[f"gain_bl:{k}"] = fnum(b.loc[k].gain_pct)
                N[f"gain_naive:{k}"] = fnum(b.loc[k].gain_naive)

    def ladder():
        m = S.ladder_macro()
        F["ladder"] = m is not None
        if m is None:
            return
        m = m[m.period == "all"]
        g = m.set_index(["date", "step"]).gain
        top = int(m.step.max())
        N["ladder_gain_end"] = fnum(g[("end", top)])
        N["ladder_gain_may"] = fnum(g[("05-01", top)])
        for s in range(1, top):
            N[f"ladder_step{s}_end"] = fnum(g[("end", s)])
        late = S.ladder_macro()
        late = late[late.period == "late"].set_index(["date", "step"]).gain
        N["ladder_gain_end_late"] = fnum(late[("end", top)])
        N["ladder_step2_end_late"] = fnum(late[("end", 2)])
        lad = S.ladder()
        e = lad[(lad.date == "end") & (lad.period == "all")].groupby("step").mae.mean()
        for s_, v in e.items():
            N[f"ladder_mae{int(s_)}"] = fnum(v, 2)
        q = lad[(lad.date == "end") & (lad.period == "all") & (lad.step == top)]
        N["ladder_gain_min"], N["ladder_gain_max"] = fnum(q.gain_vs_step0_pct.min()), fnum(q.gain_vs_step0_pct.max())

    def final():
        df = final_models()
        F["final"] = df is not None
        if df is None:
            return
        N["final_cov_min"], N["final_cov_max"] = fnum(df["cov"].min(), 2), fnum(df["cov"].max(), 2)
        end = df[df.date == "Конец сезона"]
        N["final_gain_end_min"], N["final_gain_end_max"] = fnum(end.gain.min()), fnum(end.gain.max())
        N["final_n_ens"] = str(int(df.kind.str.startswith("ens").sum()))
        N["final_n"] = str(len(df))

    def opt():
        F["opt"] = (S.TABLES / "optimization_backtest_summary.csv").exists()
        d = S._read_csv(S.TABLES / "optimization_delta_curve.csv")
        if d is not None:
            at = lambda v: d.loc[(d.delta - v).abs().idxmin()]  # noqa: E731
            N["opt_share_d5"] = fnum(at(0.05).gain_share_of_d30_pct)
            N["opt_share_d10"] = fnum(at(0.10).gain_share_of_d30_pct)
            N["opt_gain_d10_bln"] = fnum(at(0.10).gain_mln / 1e3)
            N["opt_cvar_d10_bln"] = fnum(at(0.10).cvar_gain_mln / 1e3)
            N["opt_turnover_d10"] = fnum(at(0.10).turnover_pct)
        rec = S._read_csv(S.TABLES / "optimization_recommendations.csv", dtype={"okato": str})
        if rec is not None:
            N["opt_n_regions"] = str(rec.okato.nunique())
            N["opt_pool_mha"] = fnum(rec.drop_duplicates("okato").pool_kha.sum() / 1e3, 1)
            for dv in (5, 10):
                q = rec[rec.delta_pp == dv]
                w = q.pool_kha
                m0, m1 = np.average(q.mean0, weights=w), np.average(q.mean1, weights=w)
                N[f"opt_gain_ha_d{dv}"] = fnum(m1 - m0)
                N[f"opt_gain_pct_d{dv}"] = fnum(100 * (m1 - m0) / m0)
                N[f"opt_cvar_ha_d{dv}"] = fnum(np.average(q.cvar1 - q.cvar0, weights=w))
                N[f"opt_worst_better_d{dv}"] = str(int((q.worst1 > q.worst0 + 1e-6).sum()))
                N["opt_mean0"] = fnum(m0)
        es = S._read_csv(S.TABLES / "optimization_effect_stats.csv")
        if es is not None:
            r5 = es[es.delta_pp == 5].iloc[0]
            N["opt_med_gain_d5"], N["opt_q3_gain_d5"] = fnum(r5.gain_rub_ha_median), fnum(r5.gain_rub_ha_q75)
            N["opt_med_pct_d5"], N["opt_share_win_d5"] = fnum(r5.gain_pct_median), fnum(r5.share_gain_pct)
        if d is not None and "gain_rub_per_moved_ha" in d:
            N["opt_per_moved_d5"] = fnum(round(at(0.05).gain_rub_per_moved_ha, -2))
            N["opt_gain_d5_bln"] = fnum(at(0.05).gain_mln / 1e3)
        cs = S._read_csv(S.TABLES / "optimization_constraints.csv")
        if cs is not None:
            N["opt_norot_d10_bln"] = fnum(cs[(cs.delta_pp == 10) & (cs.variant == "no_rotation")].gain_mln.iloc[0] / 1e3)
        sc = S._read_csv(S.TABLES / "optimization_scenario_compare.csv")
        if sc is not None:
            N["opt_hist_d5_bln"] = fnum(sc[(sc.delta_pp == 5) & (sc.scenarios == "история")].gain_mln.iloc[0] / 1e3)
        if rec is not None and "n_scen" in rec:
            N["opt_n_scen"] = str(int(rec.n_scen.median()))
        sens = S._read_csv(S.TABLES / "optimization_sensitivity.csv")
        if sens is not None:
            corn = sens[(sens.factor == "price") & (sens.crop_code.astype(str).str.startswith("1080202"))]
            N["opt_sens_corn_lo"], N["opt_sens_corn_hi"] = fnum(corn.gain_mln.min() / 1e3), fnum(corn.gain_mln.max() / 1e3)
            N["opt_sens_min"] = fnum(sens[sens.factor.isin(["price", "cost"])].gain_mln.min() / 1e3)
        b = S._read_csv(S.TABLES / "optimization_backtest_summary.csv")
        if b is not None:
            b = b.set_index("plan")
            N["opt_bt_n"] = fnum(b.loc["current", "n"])
            N["opt_bt_real_d10"] = fnum(b.loc["d10", "real_gain_rub_ha"])
            N["opt_bt_exp_d10"] = fnum(b.loc["d10", "exp_gain_rub_ha"])
            N["opt_bt_win_d10"] = fnum(b.loc["d10", "win_share_pct"])
            N["opt_bt_realization"] = fnum(100 * b.loc["d10", "real_gain_rub_ha"] / b.loc["d10", "exp_gain_rub_ha"])
            N["opt_bad_cur"] = fnum(b.loc["current", "bad_years_profit_rub_ha"])
            N["opt_bad_d10"] = fnum(b.loc["d10", "bad_years_profit_rub_ha"])
            N["opt_bt_bln_d10"] = fnum(b.loc["d10", "real_gain_bln_year"])

    for fn in (data, sweep, ladder, final, opt):
        block(fn)
    F["sweep_v2_report"] = S.report_exists("sweep_v2.md") or (
        S.report_exists("sweep.md") and "v2" in (S.REPORTS / "sweep.md").read_text(encoding="utf-8")[:200])
    F["opt_report"] = S.report_exists("optimization.md")
    return N, F


# ================================================================ таблицы

def _t(header, rows, widths=None, align=None, note=None):
    return {"header": header, "rows": [[str(c) for c in r] for r in rows], "widths": widths, "align": align,
            "note": note}


def t_methods(N):
    rows = [["Тренд + регрессия на погоду", "ряды урожайности, погода", "регион", "до сева и в сезоне", "да, каркас и бейзлайн"],
            ["Агрометеорологический (ГТК, влагозапасы)", "наземные агрометнаблюдения", "субъект", "весна → сезон", "идея признаков"],
            ["Процессные модели (WOFOST, DSSAT, APSIM)", "почва, сорт, агротехника, погода", "поле", "любой момент", "направление развития"],
            ["Спутниковый (NDVI и др.)", "снимки, маска культур", "поле, район", "середина вегетации", "да, блок признаков"],
            ["Бустинг и случайный лес", "панель регион × год", "регион", "до сева и в сезоне", "да, основа"],
            ["Глубокие сети (LSTM, CNN)", "тысячи наблюдений, снимки", "округ, поле", "в сезоне", "нет: мало данных"]]
    return _t(["Метод", "Данные на входе", "Уровень", "Когда даёт прогноз", "В решении"], rows, [4.2, 3.6, 2.0, 2.8, 3.0])


def t_sources(N):
    st = data_stats.load()
    desc = {"emiss": ("ЕМИСС (Росстат)", "урожайность, площади, сборы, цены, удобрения, экономика", "1986–2026"),
            "era5_om": ("ERA5 / ERA5-Land", "почасовая погода, влажность почвы, 450 точек", "2000–2026"),
            "nasa_power": ("NASA POWER", "суточная погода, 450 точек (сверка)", "2000–2026"),
            "spei": ("SPEIbase v2.9", "индекс засухи, сетка 0,5° (не вошёл: ряд до 2022)", "1901–2022"),
            "soilgrids": ("SoilGrids 2.0", "11 свойств почвы × 6 глубин, WRB", "—"),
            "modis_ndvi": ("MODIS MOD13Q1", "NDVI, 250 м, 16 дней, 450 точек", "2000–2026"),
            "modis_lai": ("MODIS MOD15A2H", "LAI и FPAR, 500 м, 8 дней", "2002–2026"),
            "modis_lst": ("MODIS MOD11A2", "температура поверхности днём и ночью, 1 км", "2000–2026"),
            "modis_gpp": ("MODIS MOD17A2H", "валовая первичная продукция", "2021–2026"),
            "crop_calendar_ggcmi": ("GGCMI Phase 3", "календари сева и созревания, 0,5°", "—"),
            "worldcover": ("ESA WorldCover", "маска пашни для выбора точек", "2021"),
            "geo": ("Границы и точки", "89 субъектов, 450 точек пашни", "—"),
            "faostat": ("FAOSTAT", "урожайность и площади стран (сверка)", "1961–2024"),
            "usda_psd": ("USDA PSD", "баланс и оценки урожая РФ", "до 2026/27"),
            "rosstat": ("Росстат (архивы)", "сборники «Регионы России», ВСХП-2016", "—"),
            "gossortrf": ("Госреестр сортов", "14 490 сортов с регионами допуска", "2026"),
            "mcx_news": ("Минсельхоз", "новости о ходе уборки", "2025–2026")}
    rows = []
    for k, (name, what, years) in desc.items():
        v = st["raw_sources"].get(k)
        if not v or v["bytes"] == 0:
            continue
        mb = v["bytes"] / 2**20
        size = f"{mb / 1024:.2f} ГБ".replace(".", ",") if mb >= 1024 else f"{mb:.1f} МБ".replace(".", ",")
        rows.append((v["bytes"], [name, what, years, size, fnum(v["files"])]))
    rows = [r for _, r in sorted(rows, key=lambda x: -x[0])]
    return _t(["Источник", "Что взято", "Годы", "Объём", "Файлов"], rows, [3.3, 6.4, 2.1, 1.9, 1.6],
              align="llll r".replace(" ", ""))


def t_sample(N):
    r = S.sweeps()
    b = S.best(r[r.ck.isin(CROP_ORDER)])
    rows = []
    for k in CROP_ORDER:
        row = b.loc[k]
        p = S.preds(row.ver, row.cfg_id)
        reg = p[p.crop == k].okato.nunique() if p is not None and "crop" in p else np.nan
        bl = {"lin10": "линейный 10 лет", "mean5": "среднее 5 лет", "lag1": "прошлый год"}
        best_bl = min(["lin10", "mean5", "lag1"], key=lambda c: row[f"mae_bl_{c}"])
        rows.append([CROP_NAME[k], fnum(reg), fnum(row.n), bl[best_bl], fnum(row.mae_best_bl, 3),
                     fnum(row.mae_bl_lag1, 3)])
    return _t(["Культура", "Регионов", "Прогнозов 2012–2025", "Лучший бейзлайн", "его MAE, т/га",
               "MAE «прошлый год», т/га"], rows, [3.6, 1.8, 2.6, 3.0, 2.2, 2.6], align="lrrlrr")


def t_clean(N):
    st = data_stats.load()
    desc = {"model_table": "субъект × культура × год: цель, тренды, все признаки",
            "yield_panel": "урожайность с флагами качества, все уровни и категории",
            "weather_region_season": "погода сезона широкой таблицей: месяцы, агрегаты, нормы",
            "weather_region_month": "погода по месяцам", "weather_crop_window": "погода в окне вегетации культуры",
            "modis_extra_region": "LAI, FPAR, LST, GPP по датам", "modis_extra_region_year": "годовые признаки MODIS",
            "ndvi_region": "NDVI по датам", "ndvi_region_year": "годовые признаки NDVI",
            "yield_operational": "оперативные данные уборки 2023–2026", "economics_aux": "удобрения, техника, себестоимость",
            "prices": "цены производителей, номинальные и реальные", "crop_calendar": "календари культур по регионам",
            "regions": "справочник субъектов", "crops": "справочник культур", "soil_region": "почвы и рельеф",
            "prices_crop_map": "соответствие продуктов и культур"}
    rows = [[f"`{k}`", desc.get(k, ""), fnum(v["rows"]), fnum(v["cols"])]
            for k, v in sorted(st["clean_tables"].items(), key=lambda kv: -kv[1]["rows"])]
    return _t(["Таблица", "Содержание", "Строк", "Колонок"], rows, [4.4, 7.6, 2.0, 1.6], align="llrr")


def t_flags(N):
    rows = [["скрытое значение (999 999 999)", "1 809", "75", "снимается"],
            ["ноль (нет данных)", "1 172", "0", "снимается"],
            ["выше трёх медиан ряда", "2 201", "41", "снимается, если не подтверждён сбором"],
            ["  — из них подтверждено сбором и площадью", "476", "27", "остаётся"],
            ["выше агрономического предела", "55", "4", "снимается"],
            ["ниже трети медианы", "5 684", "147", "остаётся, флаг"],
            ["скачок год к году в 3 раза", "7 332", "143", "остаётся, флаг"],
            ["урожайность не согласована со сбором", "9 446", "391", "остаётся; в обучении исключается"],
            ["одиночный пропуск интерполирован", "2 175", "27", "только в заполненном ряде"]]
    return _t(["Флаг", "Все строки", "6 основных культур, субъекты", "Обработка"], rows, [6.0, 2.0, 3.2, 4.4],
              align="lrrl")


def t_featsets(N):
    rows = [["F0", "только тренд и свободный член", "нижняя граница: «тренд со сдвигом»"],
            ["F1", "месячные аномалии погоды", "основной погодный сигнал"],
            ["F2", "F1 + окно культуры, зима, запас влаги", ""],
            ["F2r", "F2 + регион и прошлые отклонения", "без почв и NDVI"],
            ["F3", "F2 + климат региона и почвы", ""],
            ["F4", "F3 + NDVI", ""],
            ["F5", "F4 + регион, ФО и прошлые отклонения", ""],
            ["F6", "F5 + спутник MODIS (LAI, FPAR, LST, GPP), рельеф, WRB", "перебор v2"],
            ["F7", "спутник + рельеф + регион, без погоды", "перебор v2"],
            ["F8", "погода + окно + спутник + регион", "перебор v2, лёгкий"]]
    return _t(["Набор", "Состав", "Комментарий"], rows, [1.4, 9.0, 5.2])


def t_stages(N):
    r = S.sweeps()
    cnt = r.groupby("stage").cfg_id.nunique() if r is not None else pd.Series(dtype=int)
    desc = [("coarse", "грубый отбор", "6 трендов × 3 цели × наборы признаков × {Ridge, LightGBM}; общая модель для LightGBM"),
            ("models_default", "модели без настройки", "топ-3 сочетания культуры × 8 моделей с параметрами по умолчанию"),
            ("models", "Optuna на конце сезона", "30 попыток на исследование, подбор только по 2012–2018"),
            ("dates", "ранние даты", "лучшие модели культуры на 1 мая, 1 июня, 1 июля, 1 августа"),
            ("dates_tuned", "Optuna на каждой дате", "Ridge и LightGBM, набор признаков — гиперпараметр"),
            ("ensembles", "ансамбли", "среднее топ-2/3/5 и стекинг NNLS с весами по прошлым годам")]
    rows = [[n, what, fnum(cnt.get(k, 0))] for k, n, what in desc]
    rows.append(["разведка", "тренд × цель × признаки, проверка утечек (до перебора)", N.get("cfg_pre", "")])
    return _t(["Этап", "Что перебирается", "Конфигураций (v1 + v2)"], rows, [3.6, 9.2, 2.8], align="llr")


def _desc_cfg(row) -> str:
    m = {"ridge": "Ridge", "enet": "ElasticNet", "rf": "RandomForest", "et": "ExtraTrees", "hgb": "HistGB",
         "lgbm": "LightGBM", "xgb": "XGBoost", "cat": "CatBoost", "ens_mean2": "ансамбль топ-2",
         "ens_mean3": "ансамбль топ-3", "ens_mean5": "ансамбль топ-5", "ens_stack": "стекинг"}
    if str(row.model).startswith("ens_"):
        return m[row.model]
    tr = {"mean5": "ср. 5 лет", "lin_all": "лин.", "lin10": "лин. 10", "lin15": "лин. 15", "ts10": "ТС 10",
          "ts15": "ТС 15"}
    tg = {"diff": "разность", "ratio": "отношение", "log": "лог"}
    pooled = ", общая" if row.crop == "pooled" else ""
    return f"{m.get(row.model, row.model)}, {row.feats}, {tr.get(row.trend, row.trend)}, {tg.get(row.target, row.target)}{pooled}"


def t_best_end(N):
    r = S.sweeps()
    b = S.best(r[r.ck.isin(CROP_ORDER)])
    rows = []
    for k in CROP_ORDER:
        x = b.loc[k]
        rows.append([CROP_NAME[k], _desc_cfg(x), fnum(x.mae, 3), fnum(x.mape, 1), fnum(x.gain_pct, 1),
                     fnum(x.gain_naive, 1), fnum(x.late_gain_pct, 1), fnum(100 * x.cov80), f"±{fnum(x.width80_pct / 2)}"])
    return _t(["Культура", "Конфигурация", "MAE, т/га", "MAPE, %", "Выигрыш к бейзлайну, %", "к «прошлому году», %",
               "2019–2025, %", "Покрытие 80 %", "Интервал, %"], rows,
              [2.6, 4.6, 1.3, 1.2, 1.6, 1.6, 1.4, 1.4, 1.3], align="llrrrrrrr",
              note="Выигрыш — снижение MAE. «2019–2025» — выигрыш к бейзлайну на годах, не участвовавших в подборе гиперпараметров.")


def t_ladder(N):
    lad = S.ladder()
    if lad is None:
        return None
    q = lad[(lad.date == "end") & (lad.period == "all")]
    g = q.pivot(index="crop", columns="step", values="mae").reindex(CROP_ORDER)
    gain = q.pivot(index="crop", columns="step", values="gain_vs_step0_pct").reindex(CROP_ORDER)
    top = g.columns.max()
    rows = [[CROP_NAME[k]] + [fnum(g.loc[k, s], 2 if k == "sugar_beet" else 3) for s in g.columns]
            + [fnum(gain.loc[k, top], 1)] for k in g.index]
    names = q.drop_duplicates("step").set_index("step").step_name
    return _t(["Культура"] + [f"{s}. {names[s]}" for s in g.columns] + ["Снижение MAE 4 к 0, %"], rows,
              [3.0, 1.9, 1.9, 2.2, 2.4, 2.2, 2.0], align="l" + "r" * (len(g.columns) + 1),
              note="MAE в т/га, скользящая проверка 2012–2025, одни и те же пары «субъект × год» для всех ступеней.")


def final_models() -> pd.DataFrame | None:
    """Таблица reports/final_models.md (культура × дата → тип, MAE, выигрыш, покрытие)."""
    p = S.REPORTS / "final_models.md"
    if not p.exists():
        return None
    rows = [[c.strip() for c in ln.strip().strip("|").split("|")] for ln in p.read_text(encoding="utf-8").splitlines()
            if ln.startswith("|") and not ln.startswith("|---")]
    df = pd.DataFrame(rows[1:], columns=rows[0])
    df.columns = ["crop", "date", "kind", "members", "mae", "gain", "cov"]
    for c in ("mae", "gain", "cov"):
        df[c] = pd.to_numeric(df[c], errors="coerce")
    return df


def t_final(N):
    df = final_models()
    if df is None:
        return None
    dates = ["1 мая", "1 июня", "1 июля", "1 авг", "Конец сезона"]
    kind = {"ridge": "Ridge", "lgbm": "LightGBM", "et": "ExtraTrees", "ens_mean2": "ансамбль 2",
            "ens_mean3": "ансамбль 3", "ens_mean5": "ансамбль 5", "ens_stack": "стекинг"}
    rows = []
    for name in CROP_NAME.values():
        q = df[df.crop == name].set_index("date")
        if q.empty:
            continue
        cells = [f"{fnum(q.loc[d, 'gain'], 1)} · {fnum(q.loc[d, 'cov'], 2)}" if d in q.index else "—" for d in dates]
        end = q.loc["Конец сезона"] if "Конец сезона" in q.index else None
        rows.append([name] + cells + [kind.get(end.kind, end.kind) if end is not None else ""])
    return _t(["Культура", "1 мая", "1 июня", "1 июля", "1 августа", "Конец сезона", "Модель конца сезона"], rows,
              [3.0, 1.9, 1.9, 1.9, 1.9, 2.1, 2.4], align="lrrrrrl",
              note="В ячейке: снижение MAE к лучшему простому прогнозу, % · доля фактов внутри 80-процентного интервала. "
                   "Модель выбрана по годам 2012–2018 как самая простая в пределах 1 % от лучшей MAE.")


def t_opt_costs(N):
    c = S._read_csv(S.TABLES / "optimization_costs.csv")
    if c is None:
        return None
    rows = [[r.crop, fnum(r.cost_rub_ha_2026), f"{fnum(100 * r.max_share_sown)} %"] for r in c.itertuples()]
    return _t(["Культура", "Затраты 2026, руб./га (допущение)", "Максимальная доля в посевной площади"], rows,
              [5, 5, 5], align="lrr")


def t_opt_regions(N):
    rec = S._read_csv(S.TABLES / "optimization_recommendations.csv", dtype={"okato": str})
    if rec is None:
        return None
    q = rec[rec.delta_pp == 10].nlargest(10, "gain_mln")
    rows = [[r.name, fnum(r.pool_kha), fnum(r.gain_rub_ha), fnum(r.cvar_gain_rub_ha), fnum(r.gain_mln / 1e3, 1),
             r.change] for r in q.itertuples()]
    return _t(["Субъект", "Площадь пула, тыс. га", "Прирост E[π], руб./га", "Прирост CVaR, руб./га",
               "Эффект, млрд руб.", "Изменение плана"], rows, [3.0, 1.6, 1.7, 1.7, 1.5, 6.1], align="lrrrrl")


def t_opt_backtest(N):
    b = S._read_csv(S.TABLES / "optimization_backtest_summary.csv")
    if b is None:
        return None
    lab = {"current": "текущий план", "d5": "δ = 5 п. п.", "d10": "δ = 10 п. п.", "d20": "δ = 20 п. п.",
           "dall": "без ограничения δ", "d10_mean_only": "δ = 10, без учёта риска (λ = 0)",
           "d10_cvar_only": "δ = 10, только риск (λ = 1)", "d10_last_year": "δ = 10, по одному прошлому году"}
    rows = [[lab.get(r.plan, r.plan), fnum(r.turnover_pct, 1), fnum(r.exp_gain_rub_ha), fnum(r.real_gain_rub_ha),
             fnum(r.win_share_pct), fnum(r.bad_years_profit_rub_ha)] for r in b.itertuples()]
    return _t(["План", "Перенос площади, %", "Обещано, руб./га", "Факт, руб./га", "Выигрыш в % случаев",
               "Прибыль в плохие годы, руб./га"], rows, [4.8, 2.0, 2.0, 2.0, 2.2, 2.6], align="lrrrrr")


TABLES = {"methods": t_methods, "sources": t_sources, "sample": t_sample, "clean": t_clean, "flags": t_flags,
          "featsets": t_featsets, "stages": t_stages, "best_end": t_best_end, "ladder": t_ladder, "final": t_final,
          "opt_costs": t_opt_costs, "opt_regions": t_opt_regions, "opt_backtest": t_opt_backtest}


def keyfigures(N):
    """Плитки «Ключевые цифры»: (число, подпись)."""
    g = lambda k: N.get(k)  # noqa: E731
    tiles = [(g("n_sources"), "источников открытых данных"),
             (f"{g('raw_gb')} ГБ" if g("raw_gb") else None, f"сырых данных, {g('raw_files') or '…'} файлов"),
             (g("model_cols"), f"колонок признаков в таблице модели ({g('model_rows') or '…'} строк)"),
             (g("cfg_total"), f"конфигураций моделей, {g('optuna_trials') or '…'} попыток Optuna"),
             (f"−{g('ladder_gain_end')} %" if g("ladder_gain_end") else
              (f"−{g('gain_naive_max')} %" if g("gain_naive_max") else None),
              "ошибки к наивному прогнозу в среднем по культурам" if g("ladder_gain_end") else
              "ошибки (MAE) к наивному прогнозу «как в прошлом году», лучшая культура"),
             (f"−{g('gain_bl_max')} %" if g("gain_bl_max") else None, "ошибки к лучшему простому прогнозу (свёкла)"),
             (f"+{g('opt_bt_real_d10')} ₽/га" if g("opt_bt_real_d10") else None,
              f"фактический эффект плана посевов на истории 2012–2025, выигрыш в {g('opt_bt_win_d10') or '…'} % случаев"),
             (f"+{g('opt_gain_pct_d5')} %" if g("opt_gain_pct_d5") else None,
              f"ожидаемой прибыли при сдвиге плана не более 5 п. п. на культуру ({g('opt_n_regions') or '…'} субъектов)")]
    return tiles


# ================================================================ docx

class Doc:
    BODY_FONT = "Cambria"
    HEAD_FONT = "Segoe UI Semibold"
    ACCENT = "104281"
    BLUE = "2a78d6"

    def __init__(self, N, F, meta):
        from docx import Document
        from docx.enum.section import WD_SECTION
        from docx.shared import Cm, Pt, RGBColor
        self.Cm, self.Pt, self.RGB, self.WD_SECTION = Cm, Pt, RGBColor, WD_SECTION
        self.N, self.F, self.meta = N, F, meta
        self.d = Document()
        self.fig_n = self.tab_n = 0
        self.h1 = self.h2 = 0
        self.missing: set[str] = set()
        self._styles()

    # ---------- оформление
    def _styles(self):
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml.ns import qn
        Pt, RGB = self.Pt, self.RGB
        s = self.d.sections[0]
        s.page_height, s.page_width = self.Cm(29.7), self.Cm(21.0)
        s.left_margin, s.right_margin, s.top_margin, s.bottom_margin = (self.Cm(2.5), self.Cm(1.8),
                                                                        self.Cm(2.0), self.Cm(2.0))
        st = self.d.styles
        n = st["Normal"]
        n.font.name, n.font.size = self.BODY_FONT, Pt(10.5)
        n.element.rPr.rFonts.set(qn("w:eastAsia"), self.BODY_FONT)
        n.element.rPr.rFonts.set(qn("w:cs"), self.BODY_FONT)
        pf = n.paragraph_format
        pf.space_after, pf.line_spacing = Pt(5), 1.15
        pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        for lvl, size, color, before, after in ((1, 17, self.ACCENT, 26, 8), (2, 13, self.ACCENT, 14, 5),
                                                (3, 11.5, "2a2a28", 10, 4)):
            h = st[f"Heading {lvl}"]
            h.font.name, h.font.size, h.font.bold = self.HEAD_FONT, Pt(size), False
            h.font.color.rgb = RGB.from_string(color)
            rpr = h.element.get_or_add_rPr()
            rpr.get_or_add_rFonts().set(qn("w:ascii"), self.HEAD_FONT)
            rpr.get_or_add_rFonts().set(qn("w:hAnsi"), self.HEAD_FONT)
            rpr.get_or_add_rFonts().set(qn("w:cs"), self.HEAD_FONT)
            h.paragraph_format.space_before, h.paragraph_format.space_after = Pt(before), Pt(after)
            h.paragraph_format.keep_with_next = True
            h.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        for name in ("List Bullet", "List Number"):
            ls = st[name]
            ls.paragraph_format.space_after = Pt(2.5)
            ls.font.name = self.BODY_FONT
        from docx.enum.style import WD_STYLE_TYPE
        for lvl, size, ind in ((1, 10.5, 0), (2, 9.5, 0.6)):
            name = f"TOC {lvl}"
            ts = st[name] if name in [x.name for x in st] else st.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
            ts.base_style = st["Normal"]
            ts.font.size = Pt(size)
            ts.paragraph_format.space_after = Pt(1.5 if lvl == 2 else 3)
            ts.paragraph_format.space_before = Pt(0 if lvl == 2 else 4)
            ts.paragraph_format.left_indent = self.Cm(ind)
            ts.paragraph_format.first_line_indent = 0
            ts.paragraph_format.line_spacing = 1.0
        cap = st["Caption"]
        cap.font.name, cap.font.size, cap.font.italic, cap.font.bold = "Segoe UI", Pt(8.8), False, False
        cap.font.color.rgb = RGB.from_string("52514e")
        cap.paragraph_format.space_after = Pt(9)
        cap.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER

    def _shade(self, cell, hex_fill):
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        tcPr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), hex_fill)
        tcPr.append(shd)

    def _borders(self, table, color="d4d3cf", inside=True):
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        tblPr = table._tbl.tblPr
        b = OxmlElement("w:tblBorders")
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            e = OxmlElement(f"w:{edge}")
            if edge in ("left", "right", "insideV") or (edge == "insideH" and not inside):
                e.set(qn("w:val"), "nil")
            else:
                e.set(qn("w:val"), "single")
                e.set(qn("w:sz"), "6" if edge in ("top", "bottom") else "4")
                e.set(qn("w:color"), color)
            b.append(e)
        tblPr.append(b)

    def _field(self, run, instr):
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        for kind, text in (("begin", None), ("instr", instr), ("separate", None), ("end", None)):
            if kind == "instr":
                el = OxmlElement("w:instrText")
                el.set(qn("xml:space"), "preserve")
                el.text = text
            else:
                el = OxmlElement("w:fldChar")
                el.set(qn("w:fldCharType"), kind)
            run._r.append(el)

    # ---------- текст
    def inline(self, p, text, size=None, color=None, bold=False):
        """**жирный**, *курсив*, `код`, {{ключ}}, x_{низ}, x^{верх}."""
        from docx.enum.text import WD_COLOR_INDEX
        pat = re.compile(r"(\*\*.+?\*\*|\*[^*\s][^*]*?\*|`[^`]+`|\{\{[^}]+\}\}|_\{[^}]+\}|\^\{[^}]+\})")
        for part in pat.split(text):
            if not part:
                continue
            kw = {}
            if part.startswith("**"):
                t, kw = part[2:-2], {"bold": True}
            elif part.startswith("`"):
                t, kw = part[1:-1], {"code": True}
            elif part.startswith("{{"):
                key = part[2:-2].strip()
                t = self.N.get(key)
                if t is None:
                    self.missing.add(key)
                    r = p.add_run(f"[{PENDING}]")
                    r.font.highlight_color = WD_COLOR_INDEX.YELLOW
                    continue
            elif part.startswith("_{"):
                t, kw = part[2:-1], {"sub": True}
            elif part.startswith("^{"):
                t, kw = part[2:-1], {"sup": True}
            elif part.startswith("*") and part.endswith("*") and len(part) > 2:
                t, kw = part[1:-1], {"italic": True}
            else:
                t = part
            if "{{" in t:
                self.inline(p, t, size, color, bold)
                continue
            r = p.add_run(t)
            r.bold = bold or kw.get("bold") or None
            r.italic = kw.get("italic") or None
            if kw.get("code"):
                r.font.name = "Consolas"
                r.font.size = self.Pt(9.5)
            if kw.get("sub"):
                r.font.subscript = True
            if kw.get("sup"):
                r.font.superscript = True
            if size:
                r.font.size = self.Pt(size)
            if color:
                r.font.color.rgb = self.RGB.from_string(color)

    def heading(self, level, text):
        numbered = not re.match(r"(Аннотация|Заключение|Список источников|Ключевые цифры)", text)
        if level == 1:
            if text.startswith("Список источников"):
                self.page_break()
            if numbered:
                self.h1 += 1
                self.h2 = 0
                text = f"{self.h1}. {text}"
        elif level == 2 and self.h1 and numbered and self._cur_numbered:
            self.h2 += 1
            text = f"{self.h1}.{self.h2}. {text}"
        if level == 1:
            self._cur_numbered = numbered
        self.d.add_heading(text, level=level)

    _cur_numbered = False

    def para(self, text, style=None):
        p = self.d.add_paragraph(style=style)
        if style is None:
            p.paragraph_format.first_line_indent = self.Cm(0.75)
        self.inline(p, text)
        return p

    def bullet(self, text, numbered=False, num=None):
        if numbered:  # нумерацию ведём сами: списки Word не перезапускаются по разделам
            p = self.d.add_paragraph()
            p.paragraph_format.left_indent = self.Cm(0.9)
            p.paragraph_format.first_line_indent = self.Cm(-0.5)
            p.paragraph_format.space_after = self.Pt(2.5)
            self.inline(p, f"{num}.  " + text)
        else:
            p = self.d.add_paragraph(style="List Bullet")
            self.inline(p, text)
        return p

    def formula(self, text):
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        p = self.d.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = self.Pt(4)
        p.paragraph_format.space_after = self.Pt(8)
        for part in re.split(r"(_\{[^}]+\}|\^\{[^}]+\})", text):
            if not part:
                continue
            sub, sup = part.startswith("_{"), part.startswith("^{")
            r = p.add_run(part[2:-1] if (sub or sup) else part)
            r.font.name, r.font.size, r.italic = "Cambria Math", self.Pt(11.5), True
            r.font.subscript = sub or None
            r.font.superscript = sup or None

    def page_break(self):
        from docx.enum.text import WD_BREAK
        self.d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def caption(self, kind, text):
        p = self.d.add_paragraph(style="Caption")
        if kind == "Таблица":
            self.tab_n += 1
            from docx.enum.text import WD_ALIGN_PARAGRAPH
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.space_before = self.Pt(6)
            p.paragraph_format.space_after = self.Pt(3)
            r = p.add_run(f"Таблица {self.tab_n}. ")
        else:
            self.fig_n += 1
            r = p.add_run(f"Рисунок {self.fig_n}. ")
        r.bold = True
        self.inline(p, text)

    def pending(self, text):
        from docx.enum.text import WD_COLOR_INDEX
        p = self.d.add_paragraph()
        r = p.add_run(f"[{PENDING}] ")
        r.bold = True
        r.font.highlight_color = WD_COLOR_INDEX.YELLOW
        self.inline(p, text)
        for run in p.runs:
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW
        self.missing.add("pending: " + text[:60])

    def figure(self, path: Path | None, caption, width_cm=16.2, max_h=9.2):
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        if path is None or not path.exists():
            self.pending(f"Рисунок «{caption}» появится после расчёта.")
            return
        p = self.d.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.space_before = self.Pt(6)
        p.paragraph_format.space_after = self.Pt(2)
        from PIL import Image
        with Image.open(path) as im:
            w, h = im.size
            dpi = (im.info.get("dpi") or (220, 220))[0] or 220
        width_cm = min(width_cm, w / dpi * 2.54)
        if width_cm * h / w > max_h:
            width_cm = max_h * w / h
        p.add_run().add_picture(str(path), width=self.Cm(width_cm))
        self.caption("Рисунок", caption)

    def table(self, spec, caption):
        if spec is None:
            self.pending(f"Таблица «{caption}» появится после расчёта.")
            return
        from docx.enum.table import WD_TABLE_ALIGNMENT
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        self.caption("Таблица", caption)
        hdr, rows = spec["header"], spec["rows"]
        t = self.d.add_table(rows=1 + len(rows), cols=len(hdr))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        self._borders(t)
        align = spec.get("align") or "l" * len(hdr)
        small = len(hdr) >= 6 or len(rows) > 12
        fs = 8 if small else 8.8
        for j, h in enumerate(hdr):
            c = t.rows[0].cells[j]
            self._shade(c, "e6f0fc")
            c.paragraphs[0].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT if align[j] == "l" else WD_ALIGN_PARAGRAPH.RIGHT
            c.paragraphs[0].paragraph_format.first_line_indent = 0
            self.inline(c.paragraphs[0], h, size=fs, bold=True, color="0d366b")
        for i, row in enumerate(rows, start=1):
            for j, v in enumerate(row):
                c = t.rows[i].cells[j]
                para = c.paragraphs[0]
                para.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT if align[j] == "l" else WD_ALIGN_PARAGRAPH.RIGHT
                para.paragraph_format.space_after = self.Pt(1)
                para.paragraph_format.line_spacing = 1.0
                self.inline(para, v, size=fs)
        keep = len(rows) <= 14  # короткую таблицу не разрываем между страницами
        for ri, row in enumerate(t.rows):
            for c in row.cells:
                c.paragraphs[0].paragraph_format.space_before = self.Pt(1.5)
                c.paragraphs[0].paragraph_format.space_after = self.Pt(1.5)
                c.paragraphs[0].paragraph_format.keep_with_next = keep and ri < len(t.rows) - 1
        if spec.get("widths"):
            tot = sum(spec["widths"])
            for j, w in enumerate(spec["widths"]):
                for row in t.rows:
                    row.cells[j].width = self.Cm(16.6 * w / tot)
        # строка заголовка повторяется на новой странице
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        trPr = t.rows[0]._tr.get_or_add_trPr()
        th = OxmlElement("w:tblHeader")
        th.set(qn("w:val"), "true")
        trPr.append(th)
        if spec.get("note"):
            p = self.d.add_paragraph()
            p.paragraph_format.space_before = self.Pt(2)
            self.inline(p, spec["note"], size=8.3, color="52514e")
        else:
            self.d.add_paragraph().paragraph_format.space_after = self.Pt(2)

    def md_table(self, lines, caption=None):
        rows = [[c.strip() for c in ln.strip().strip("|").split("|")] for ln in lines]
        rows = [r for r in rows if not all(re.fullmatch(r":?-+:?", c) for c in r)]
        hdr, body = rows[0], rows[1:]
        align = "".join("r" if all(re.fullmatch(r"[−\-+±]?[\d\s ,.%]+", r[j] or "x") for r in body) else "l"
                        for j in range(len(hdr)))
        self.tab_n -= 0
        self.table(_t(hdr, body, align=align), caption or "")

    def keyfigures(self, tiles):
        from docx.enum.table import WD_TABLE_ALIGNMENT
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        h = self.d.add_heading("Ключевые цифры", level=2)
        h.paragraph_format.space_before = self.Pt(14)
        t = self.d.add_table(rows=2, cols=4)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        self._borders(t, color="ffffff", inside=False)
        for i, (val, cap) in enumerate(tiles):
            c = t.rows[i // 4].cells[i % 4]
            c.width = self.Cm(4.15)
            self._shade(c, "f1f6fd")
            p = c.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.first_line_indent = 0
            p.paragraph_format.space_before = self.Pt(6)
            p.paragraph_format.space_after = self.Pt(1)
            if val:
                r = p.add_run(str(val))
            else:
                from docx.enum.text import WD_COLOR_INDEX
                r = p.add_run(PENDING)
                r.font.highlight_color = WD_COLOR_INDEX.YELLOW
                self.missing.add("keyfigure: " + cap[:40])
            r.font.name, r.font.size = self.HEAD_FONT, self.Pt(19 if val and len(str(val)) <= 9 else 15)
            r.font.color.rgb = self.RGB.from_string(self.BLUE)
            p2 = c.add_paragraph()
            p2.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p2.paragraph_format.first_line_indent = 0
            p2.paragraph_format.space_after = self.Pt(7)
            p2.paragraph_format.line_spacing = 1.0
            self.inline(p2, cap, size=8.4, color="3d3c39")
        self.d.add_paragraph()

    # ---------- титул, оглавление, колонтитулы
    def title_page(self):
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        m = self.meta
        for _ in range(3):
            self.d.add_paragraph()

        def line(text, size, color="0b0b0b", font=None, bold=False, after=6):
            p = self.d.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_after = self.Pt(after)
            r = p.add_run(text)
            r.font.size, r.bold = self.Pt(size), bold
            r.font.name = font or self.HEAD_FONT
            r.font.color.rgb = self.RGB.from_string(color)
            return p

        line(m["case"], 10, "52514e", after=28)
        line(m["kind"].upper(), 10.5, self.BLUE, after=4)
        line(m["title"], 28, self.ACCENT, after=10)
        line(m["subtitle"], 13, "3d3c39", font=self.BODY_FONT, after=40)
        for _ in range(12):
            self.d.add_paragraph()
        author = m.get("author") or "[ФИО автора]"
        line(author, 12, "0b0b0b", bold=True, after=2)
        line(m["affiliation"], 10, "52514e", font=self.BODY_FONT, after=2)
        line(f"{m['city_year']} · версия от {self.N['date_build']}", 10, "52514e", font=self.BODY_FONT)
        if m.get("subtitle_links"):
            line(m["subtitle_links"], 9.5, self.BLUE, font=self.BODY_FONT)
        self.page_break()

    def toc(self):
        h = self.d.add_paragraph()
        r = h.add_run("Содержание")
        r.font.name, r.font.size = self.HEAD_FONT, self.Pt(17)
        r.font.color.rgb = self.RGB.from_string(self.ACCENT)
        p = self.d.add_paragraph()
        self._field(p.add_run(), 'TOC \\o "1-2" \\h \\z \\u')
        p.add_run("Оглавление обновляется при сборке PDF (или в Word: F9).")
        self.page_break()

    def headers(self):
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        s = self.d.sections[0]
        s.different_first_page_header_footer = True
        hp = s.header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r = hp.add_run(f"{self.meta['title']} · научный отчёт")
        r.font.size, r.font.name = self.Pt(8), "Segoe UI"
        r.font.color.rgb = self.RGB.from_string("8a8984")
        fp = s.footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = fp.add_run()
        run.font.size, run.font.name = self.Pt(9), "Segoe UI"
        self._field(run, "PAGE")

    def save(self, path: Path):
        path.parent.mkdir(parents=True, exist_ok=True)
        self.d.core_properties.title = self.meta["title"]
        self.d.core_properties.subject = self.meta["subtitle"]
        self.d.core_properties.author = self.meta.get("author") or ""
        self.d.save(path)


# ================================================================ разбор markdown

def render(doc: Doc, md: str):
    lines = md.splitlines()
    stack: list[bool] = []  # активность ветвей ::if
    i = 0
    para: list[str] = []
    numctr = 0

    def active():
        return all(stack)

    def flush():
        nonlocal para
        if para:
            doc.para(" ".join(s.strip() for s in para))
            para = []

    while i < len(lines):
        ln = lines[i].rstrip()
        s = ln.strip()
        if s.startswith("::if "):
            flush()
            stack.append(bool(doc.F.get(s[5:].strip())))
            i += 1
            continue
        if s == "::else":
            flush()
            stack[-1] = not stack[-1]
            i += 1
            continue
        if s == "::endif":
            flush()
            stack.pop()
            i += 1
            continue
        if not active():
            i += 1
            continue
        if not s:
            flush()
            numctr = 0 if not (i + 1 < len(lines) and re.match(r"\d+\. ", lines[i + 1].strip())) else numctr
            i += 1
            continue
        m = re.match(r"^(#{1,3})\s+(.*)", s)
        if m:
            flush()
            doc.heading(len(m.group(1)), m.group(2))
            numctr = 0
            i += 1
            continue
        if s.startswith("::"):
            flush()
            name, _, arg = s[2:].partition(" ")
            a, _, cap = arg.partition("|")
            a, cap = a.strip(), cap.strip()
            if name == "figure":
                doc.figure(FIG / f"{a}.png" if (FIG / f"{a}.png").exists() else None, cap)
            elif name == "image":
                doc.figure(ROOT / a if (ROOT / a).exists() else None, cap, width_cm=15.5, max_h=7.8)
            elif name == "table":
                fn = TABLES.get(a)
                try:
                    spec = fn(doc.N) if fn else None
                except Exception as e:  # noqa: BLE001
                    print(f"[table] {a}: {e!r}", file=sys.stderr)
                    spec = None
                doc.table(spec, cap)
            elif name == "keyfigures":
                doc.keyfigures(keyfigures(doc.N))
            elif name == "pending":
                doc.pending(arg.strip())
            elif name == "pagebreak":
                doc.page_break()
            i += 1
            continue
        if s.startswith("$$"):
            flush()
            doc.formula(s.strip("$").strip())
            i += 1
            continue
        if s.startswith("|"):
            flush()
            block = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                block.append(lines[i])
                i += 1
            doc.md_table(block)
            continue
        if re.match(r"^[-*] ", s):
            flush()
            doc.bullet(s[2:])
            i += 1
            continue
        m = re.match(r"^(\d+)\. (.*)", s)
        if m:
            flush()
            numctr += 1
            doc.bullet(m.group(2), numbered=True, num=m.group(1))
            i += 1
            continue
        para.append(s)
        i += 1
    flush()


def references(doc: Doc, md: str):
    """Список источников: нумерация из файла, висячий отступ."""
    for ln in md.splitlines():
        s = ln.strip()
        if s.startswith("# "):
            doc.heading(1, s[2:])
        elif s.startswith("## "):
            doc.d.add_heading(s[3:], level=2)
        elif m := re.match(r"^(\d+)\. (.*)", s):
            p = doc.d.add_paragraph()
            p.paragraph_format.left_indent = doc.Cm(0.8)
            p.paragraph_format.first_line_indent = doc.Cm(-0.8)
            p.paragraph_format.space_after = doc.Pt(3)
            from docx.enum.text import WD_ALIGN_PARAGRAPH
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            doc.inline(p, f"{m.group(1)}.  {m.group(2)}", size=9.5)


# ================================================================ PDF через Word

def word_finalize(docx: Path, pdf: Path | None) -> int | None:
    """Обновляет оглавление и поля в Word, сохраняет .docx и экспортирует PDF. Возвращает число страниц."""
    try:
        import pythoncom
        import win32com.client as win32
    except ImportError:
        print("pywin32 недоступен: оглавление не обновлено, PDF не собран", file=sys.stderr)
        return None
    pythoncom.CoInitialize()
    word = win32.DispatchEx("Word.Application")
    word.Visible = False
    word.DisplayAlerts = 0
    pages = None
    try:
        d = word.Documents.Open(str(docx.resolve()))
        for t in d.TablesOfContents:
            t.Update()
        d.Fields.Update()
        for t in d.TablesOfContents:
            t.Update()
        pages = d.ComputeStatistics(2)
        d.Save()
        if pdf is not None:
            d.ExportAsFixedFormat(str(pdf.resolve()), 17, OptimizeFor=0, CreateBookmarks=1, DocStructureTags=True)
        d.Close(False)
    finally:
        word.Quit()
        pythoncom.CoUninitialize()
    return pages


# ================================================================ main

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--no-pdf", action="store_true", help="не запускать Word: только .docx без обновления оглавления")
    ap.add_argument("--stats", action="store_true", help="пересчитать инвентарь data/ (≈10 мин)")
    ap.add_argument("--shap", action="store_true", help="пересчитать SHAP (≈2 мин)")
    a = ap.parse_args()
    if a.stats:
        data_stats.load(force=True)
    if a.shap or not (S.BUILD / "shap_features.csv").exists():
        from report import shap_calc
        shap_calc.main()
    style.setup()
    FIG.mkdir(parents=True, exist_ok=True)
    for k, fn in figures.FIGS.items():
        p = FIG / f"{k}.png"
        try:
            ok = fn(p)
        except Exception as e:  # noqa: BLE001
            print(f"[figure] {k}: {e!r}", file=sys.stderr)
            ok = False
        if not ok and p.exists():
            p.unlink()
    N, F = compute_numbers()
    meta = tomllib.loads((ROOT / "report" / "meta.toml").read_text(encoding="utf-8"))
    doc = Doc(N, F, meta)
    doc.headers()
    doc.title_page()
    files = sorted(TEXT.glob("*.md"))
    for f in files:
        md = f.read_text(encoding="utf-8")
        if f.name.startswith("00_"):
            render(doc, md)
            doc.page_break()
            doc.toc()
            continue
        if f.name.startswith("99_"):
            references(doc, md)
        else:
            render(doc, md)
    out_docx = OUT / f"{NAME}.docx"
    doc.save(out_docx)
    pages = None
    if not a.no_pdf:
        try:
            pages = word_finalize(out_docx, OUT / f"{NAME}.pdf")
        except Exception as e:  # noqa: BLE001
            print(f"PDF не собран: {e!r}", file=sys.stderr)
    print(f"готово: {out_docx}" + (f", страниц {pages}" if pages else ""))
    print(f"рисунков {doc.fig_n}, таблиц {doc.tab_n}")
    if doc.missing:
        print("ждут данных:", "; ".join(sorted(doc.missing)))
    (S.BUILD / "missing.txt").write_text("\n".join(sorted(doc.missing)), encoding="utf-8")


if __name__ == "__main__":
    main()
